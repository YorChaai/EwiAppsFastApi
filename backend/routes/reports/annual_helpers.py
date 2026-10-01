import logging
import re

from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# Setup logging
logger = logging.getLogger(__name__)

# Column width constants
COLUMN_WIDTH_CURRENCY = 5
COLUMN_WIDTH_EXCHANGE = 10
COLUMN_WIDTH_CATEGORY = 18
ROW_HEIGHT_HEADER_VERTICAL = 60
ROW_HEIGHT_DIVIDEND = 20

# Definisi Lebar Kolom Minimum (MIN_WIDTHS)
MIN_WIDTHS = {
    "A": 0.6328125,
    "B": 11.81640625,
    "C": 4.0,
    "D": 56.26953125,
    "E": 15.0,  # Widened from 4.63 to fit INVOICE VALUE
    "F": 13.36328125,
    "G": 5.0,
    "H": 10.0,
    "I": 18.0,
    "J": 18.0,
    "K": 18.0,
    "L": 18.0,
    "M": 18.0,
    "N": 18.0,
    "O": 18.0,
    "P": 18.0,
    "Q": 18.0,
    "R": 18.0,  # Added to ensure Pemasukan Negara fits
    "S": 18.0,  # Added to ensure Pemasukan Negara fits
}

# Row constants
EXPENSE_HEADER_ROW = 40
EXPENSE_START_ROW = 41
REVENUE_START_ROW = 8
TAX_START_ROW = 26

# Template row limits
REVENUE_TEMPLATE_END = 22
TAX_TEMPLATE_END = 37

PARENT_CATEGORY_NAMES = {
    "biaya operasi",
    "biaya research (r&d)",
    "biaya sewa peralatan",
    "biaya interpretasi log data",
    "administrasi",
    "pembelian barang",
    "sewa kantor",
    "kesehatan",
    "bisnis dev",
}

# Styles
GREEN_FILL = PatternFill(fill_type="solid", fgColor="C6EFCE")
BLUE_FILL = PatternFill(fill_type="solid", fgColor="92CDDC")
THIN_BORDER = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)


def _extract_subcategory_from_description(description):
    """Ekstrak sub-kategori dari deskripsi dengan format [SubCategory] Description"""
    if not description:
        return ""
    match = re.match(r"^\[(.*?)\]\s*", description.strip())
    if match:
        return match.group(1).strip()
    return ""


def _root_category_info(category_id, category_by_id):
    cat = category_by_id.get(category_id)
    if not cat:
        return "-", "-"
    subcategory_name = cat.name
    while cat.parent_id and category_by_id.get(cat.parent_id):
        cat = category_by_id[cat.parent_id]
    return cat.name or "-", subcategory_name or "-"


def _subcategory_sort_key(label):
    from utils.report_helpers import _safe_text

    text = _safe_text(label).strip()
    if not text or text == "-":
        return (1, "")
    return (0, text.lower())


def _clear_merged_ranges_in_region(ws, start_row, end_row, start_col, end_col):
    """Find and unmerge all ranges in the specified region."""
    ranges_to_unmerge = []
    for merged_range in list(ws.merged_cells.ranges):
        min_col, min_row, max_col, max_row = merged_range.bounds
        overlaps = not (
            max_row < start_row
            or min_row > end_row
            or max_col < start_col
            or min_col > end_col
        )
        if overlaps:
            ranges_to_unmerge.append(merged_range)

    for r_range in ranges_to_unmerge:
        try:
            ws.unmerge_cells(r_range.coord)
        except Exception:
            pass


def _write_dynamic_category_headers(
    ws, root_cats, header_row=40, start_col=9, template_end_col=30
):
    no_fill = PatternFill(fill_type=None)
    no_border = Border(
        left=Side(style=None),
        right=Side(style=None),
        top=Side(style=None),
        bottom=Side(style=None),
    )

    last_cat_col = start_col + len(root_cats) - 1

    # 1. Clear any old extra columns up to template_end_col
    for col in range(start_col, max(template_end_col + 1, last_cat_col + 10)):
        cell = ws.cell(row=header_row, column=col)
        if isinstance(cell, MergedCell):
            continue
        cell.value = None
        cell.fill = no_fill
        cell.border = no_border

    # 2. Write extra headers immediately after the last category
    extra_headers = ["Sub total", "Total", "From", "To"]

    # Get reference fill from Date column (col 2)
    ref_cell = ws.cell(row=header_row, column=2)
    theme6_fill = (
        ref_cell.fill.copy()
        if (ref_cell.fill and ref_cell.fill.fill_type)
        else PatternFill()
    )

    for idx, header_text in enumerate(extra_headers):
        col = last_cat_col + 1 + idx
        cell = ws.cell(row=header_row, column=col)
        if isinstance(cell, MergedCell):
            continue
        cell.value = header_text
        cell.font = Font(bold=True, size=12, name="Arial Narrow")
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )
        cell.fill = theme6_fill

    for offset, category in enumerate(root_cats):
        col = start_col + offset
        cell = ws.cell(row=header_row, column=col)
        if isinstance(cell, MergedCell):
            continue
        cell.value = category.name
        cell.fill = GREEN_FILL
        cell.font = Font(bold=True, size=12, name="Arial Narrow")
        cell.border = THIN_BORDER
        cell.alignment = Alignment(
            wrap_text=True, horizontal="center", vertical="center"
        )


def _set_dividend_row_border(ws, row):
    """Set border for dividend row with NO vertical line between D and E."""
    ws.cell(row=row, column=2).border = THIN_BORDER
    ws.cell(row=row, column=3).border = THIN_BORDER
    ws.cell(row=row, column=4).border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style=None),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )
    ws.cell(row=row, column=5).border = Border(
        left=Side(style=None),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )


def _compute_dividend_distribution_v2(revenues, expenses, setting, recipient_count):
    revenue_direct = sum(
        (r.idr_invoice_value or 0)
        for r in revenues
        if (r.revenue_type or "pendapatan_langsung") == "pendapatan_langsung"
    )
    revenue_other_db = sum(
        (r.idr_invoice_value or 0)
        for r in revenues
        if r.revenue_type == "pendapatan_lain_lain"
    )

    # In Excel Summary: Revenue = Direct + Other (DB)
    revenue_summary = revenue_direct + revenue_other_db

    pph23_total = sum(
        (r.pph_23 or 0.0)
        for r in revenues
    )

    # ? FIX: Filter out "TOTAL" or "-" rows which might have been imported as trash expenses
    total_cost = sum(
        (e.idr_amount or 0.0)
        for e in expenses
        if str(e.description or "").strip().upper()
        not in ["TOTAL", "-", "", "SUBTOTAL"]
    )

    other_income_setting = (
        getattr(setting, "other_income", 0.0) if setting else 0.0
    ) or 0.0

    # Profit Sebelum Tax = Revenue Summary - Total Cost + Other Income (Setting)
    profit_before_tax = revenue_summary - total_cost + other_income_setting

    # Auto calc for summary: PPh Badan (using corporate_tax_rate from setting, default 11%)
    tax_rate_raw = (
        getattr(setting, "corporate_tax_rate", 0.11)
        if setting and getattr(setting, "corporate_tax_rate", None) is not None
        else 0.11
    ) or 0.11

    # ? FIX: Handle percentage vs decimal (e.g. 22.0 vs 0.22)
    tax_rate = tax_rate_raw / 100.0 if tax_rate_raw > 1.0 else tax_rate_raw

    tax_corporate_auto = profit_before_tax * tax_rate if profit_before_tax > 0 else 0.0

    # Manual input for Table 7
    tax_pph29_input = (
        getattr(setting, "corporate_tax_pph29", 0.0) if setting else 0.0
    ) or 0.0

    # Tax Shortfall (Kekurangan Bayar) = PPh Badan - PPh 23
    tax_shortfall = max(tax_corporate_auto - pph23_total, 0.0)

    # Profit After Tax (Accounting/Summary) = Profit Sebelum Tax - Kekurangan Bayar
    # Note: Using max to avoid negative shortfall adding to profit
    profit_after_tax_accounting = profit_before_tax - tax_shortfall

    # Profit After Tax (Cash/Table 7) = Profit Sebelum Tax - Tax Input
    profit_after_tax_cash = profit_before_tax - tax_pph29_input

    profit_retained_setting = (
        getattr(setting, "profit_retained", 0.0) if setting else 0.0
    ) or 0.0
    dividend_distributed = max(
        profit_after_tax_accounting - profit_retained_setting, 0.0
    )
    dividend_per_person = (
        dividend_distributed / recipient_count if recipient_count > 0 else 0.0
    )

    return {
        "revenue_total": revenue_summary,  # This matches Excel's "Revenue" in summary
        "revenue_direct": revenue_direct,
        "revenue_other_db": revenue_other_db,
        "other_income_setting": other_income_setting,
        "pph23_total": pph23_total,
        "total_cost": total_cost,
        "profit_before_tax": profit_before_tax,
        "tax_corporate_auto": tax_corporate_auto,
        "tax_shortfall": tax_shortfall,
        "tax_pph29_input": tax_pph29_input,
        "profit_after_tax_accounting": profit_after_tax_accounting,
        "profit_after_tax_cash": profit_after_tax_cash,
        "profit_retained_setting": profit_retained_setting,
        "dividend_distributed": dividend_distributed,
        "dividend_per_person": dividend_per_person,
    }
