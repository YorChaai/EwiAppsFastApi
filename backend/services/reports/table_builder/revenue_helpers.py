import logging
import re
from datetime import datetime
from io import BytesIO
from typing import Dict, List

logger = logging.getLogger(__name__)
from models import (
    Category,
    ManualCombineGroup,
)
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from services.reports.annual_service.annual_constants import (
    COL_AMOUNT_RECEIVED_IDR,
    COL_AMOUNT_RECEIVED_USD,
    COL_CLIENT,
    COL_CURRENCY,
    COL_EXCHANGE_RATE,
    COL_INVOICE_NUMBER,
    COL_INVOICE_VALUE_IDR,
    COL_INVOICE_VALUE_USD,
    COL_PPH_23,
    COL_PPN,
    COL_RECEIVE_DATE,
    COL_REMARK,
    COL_TAX_CURRENCY,
    COL_TAX_DATE,
    COL_TAX_DPP_PEMASUKAN_NEGARA,
    COL_TAX_DPP_PPH21,
    COL_TAX_DPP_PPH23,
    COL_TAX_DPP_PPH26,
    COL_TAX_DPP_PPN,
    COL_TAX_INV_VALUE_USD,
    COL_TAX_RATE,
    COL_TAX_TRANS_VALUE,
    COL_TAX_TRANS_VALUE_IDR,
    COL_TAX_VALUE_PEMASUKAN_NEGARA,
    COL_TAX_VALUE_PPH21,
    COL_TAX_VALUE_PPH23,
    COL_TAX_VALUE_PPH26,
    COL_TAX_VALUE_PPN,
    COL_TRANSFER_FEE,
)
from services.reports.annual_service.annual_helpers import (
    BLUE_FILL,
    GREEN_FILL,
    REVENUE_START_ROW,
    REVENUE_TEMPLATE_END,
    THIN_BORDER,
    _clear_merged_ranges_in_region,
    _root_category_info,
)
from services.reports.table_builder.expense_renderer_utils import (
    _expense_amount_for_display,
    _expense_column_mapping_name,
)

from utils.report_helpers import (
    _as_iso_date,
    _clean_currency,
    _clean_settlement_title,
    _clear_range,
    _clear_range_force,
    _extract_imported_row,
    _get_dynamic_row_height,
    _idr_from_currency,
    _is_batch_settlement,
    _map_expense_category_index_from_name,
    _map_expense_column,
    _parse_iso_date,
    _safe_set_cell,
    _safe_set_number,
    _safe_text,
    _set_date_with_format,
    _set_formula_with_format,
    _set_rows_hidden,
    _shorten,
    _to_float,
)

def _manual_combine_groups_by_table(table_name, year, company="EWI Foundation"):
    groups = (
        ManualCombineGroup.query.filter_by(
            table_name=table_name,
            report_year=int(year),
            company=company,
        )
        .order_by(ManualCombineGroup.group_date.asc(), ManualCombineGroup.id.asc())
        .all()
    )
    return [group.to_dict() for group in groups]

def _manual_combine_groups_by_rows(table_name, row_ids, company="EWI Foundation"):
    """
    Ambil combine groups yang benar-benar relevan dengan data baris saat ini.
    Penting untuk mode range/all-year agar tidak tergantung report_year default.
    """
    try:
        row_id_set = {int(rid) for rid in (row_ids or [])}
    except Exception:
        row_id_set = set()

    if not row_id_set:
        return []

    groups = (
        ManualCombineGroup.query.filter_by(table_name=table_name, company=company)
        .order_by(ManualCombineGroup.group_date.asc(), ManualCombineGroup.id.asc())
        .all()
    )

    matched = []
    for group in groups:
        gids = set(group.row_ids())
        if gids.intersection(row_id_set):
            matched.append(group.to_dict())
    return matched

def _apply_manual_revenue_combine_groups(ws, revenues, combine_groups, start_row=8):
    from services.reports.annual_service.annual_excel import _set_merged_top_alignment

    if not revenues:
        return

    if not combine_groups:
        # Auto-merge consecutive rows with the same invoice_number
        combine_groups = []
        current_group = []
        last_inv = None
        for r in revenues:
            inv = r.get("invoice_number")
            rid = r.get("id")
            if inv and str(inv).strip() and inv == last_inv:
                current_group.append(rid)
            else:
                if len(current_group) > 1:
                    combine_groups.append({"row_ids": current_group})
                current_group = [rid] if (inv and str(inv).strip()) else []
                last_inv = inv
        if len(current_group) > 1:
            combine_groups.append({"row_ids": current_group})

    if not combine_groups:
        return

    row_by_id = {
        int(revenue.get("id") or 0): start_row + idx
        for idx, revenue in enumerate(revenues)
        if revenue.get("id") is not None
    }

    # Map field name to Excel Column Index
    field_to_col = {
        "invoice_date": 2,  # COL_INVOICE_DATE
        "description": 4,  # COL_DESCRIPTION
        "invoice_value": COL_INVOICE_VALUE_IDR,  # Fallback
        "invoice_value_usd": COL_INVOICE_VALUE_USD,
        "invoice_value_idr": COL_INVOICE_VALUE_IDR,
        "currency": COL_CURRENCY,
        "currency_exchange": COL_EXCHANGE_RATE,
        "invoice_number": COL_INVOICE_NUMBER,
        "client": COL_CLIENT,
        "receive_date": COL_RECEIVE_DATE,
        "amount_received": COL_AMOUNT_RECEIVED_IDR,  # Fallback
        "amount_received_idr": COL_AMOUNT_RECEIVED_IDR,
        "amount_received_usd": COL_AMOUNT_RECEIVED_USD,
        "ppn": COL_PPN,
        "pph_23": COL_PPH_23,
        "transfer_fee": COL_TRANSFER_FEE,
        "remark": COL_REMARK,
    }

    for group in combine_groups:
        group_ids = [
            int(rid) for rid in group.get("row_ids", []) if int(rid) in row_by_id
        ]
        if len(group_ids) < 2:
            continue

        row_numbers = [row_by_id[rid] for rid in group_ids]
        row_start, row_end = min(row_numbers), max(row_numbers)

        merged_cols = group.get("merged_columns") or []
        # Fallback for old groups without merged_columns info
        if not merged_cols:
            merged_cols = [
                "receive_date",
                "amount_received",
                "amount_received_idr",
                "amount_received_usd",
                "ppn",
                "pph_23",
            ]

        for field in merged_cols:
            col_idx = field_to_col.get(field)
            if not col_idx:
                continue

            # Perform Excel Merge
            try:
                # Safely unmerge if there's any existing overlapping merge
                ranges_to_unmerge = []
                for merged_range in list(ws.merged_cells.ranges):
                    min_col, min_row, max_col, max_row = merged_range.bounds
                    if not (
                        max_row < row_start
                        or min_row > row_end
                        or max_col < col_idx
                        or min_col > col_idx
                    ):
                        ranges_to_unmerge.append(merged_range)

                for r_range in ranges_to_unmerge:
                    try:
                        ws.unmerge_cells(str(r_range))
                    except Exception:
                        pass

                ws.merge_cells(
                    start_row=row_start,
                    start_column=col_idx,
                    end_row=row_end,
                    end_column=col_idx,
                )
                _set_merged_top_alignment(ws, row_start, col_idx)
            except Exception:
                pass

def _apply_manual_tax_combine_groups(ws, taxes, combine_groups, start_row=27):
    from services.reports.annual_service.annual_excel import _set_merged_top_alignment

    if not taxes:
        return

    if not combine_groups:
        combine_groups = []
        current_group = []
        last_key = None
        for t in taxes:
            key = (t.get("date"), t.get("transaction_value_idr"))
            rid = t.get("id")
            if key[0] and key[1] and key == last_key:
                current_group.append(rid)
            else:
                if len(current_group) > 1:
                    combine_groups.append({"row_ids": current_group})
                current_group = [rid] if (key[0] and key[1]) else []
                last_key = key
        if len(current_group) > 1:
            combine_groups.append({"row_ids": current_group})

    if not combine_groups:
        return

    row_by_id = {
        int(item.get("id") or 0): start_row + idx
        for idx, item in enumerate(taxes)
        if item.get("id") is not None
    }

    field_to_col = {
        "date": COL_TAX_DATE,
        "transaction_value": COL_TAX_TRANS_VALUE,
        "invoice_value_usd": COL_TAX_INV_VALUE_USD,
        "currency": COL_TAX_CURRENCY,
        "currency_exchange": COL_TAX_RATE,
        "transaction_value_idr": COL_TAX_TRANS_VALUE_IDR,
        "ppn_base": COL_TAX_DPP_PPN,
        "ppn": COL_TAX_VALUE_PPN,
        "pph_21_base": COL_TAX_DPP_PPH21,
        "pph_21": COL_TAX_VALUE_PPH21,
        "pph_23_base": COL_TAX_DPP_PPH23,
        "pph_23": COL_TAX_VALUE_PPH23,
        "pph_26_base": COL_TAX_DPP_PPH26,
        "pph_26": COL_TAX_VALUE_PPH26,
        "pemasukan_negara_base": COL_TAX_DPP_PEMASUKAN_NEGARA,
        "pemasukan_negara": COL_TAX_VALUE_PEMASUKAN_NEGARA,
    }

    for group in combine_groups:
        group_ids = [
            int(rid) for rid in group.get("row_ids", []) if int(rid) in row_by_id
        ]
        if len(group_ids) < 2:
            continue

        row_numbers = [row_by_id[rid] for rid in group_ids]
        row_start, row_end = min(row_numbers), max(row_numbers)

        merged_cols = group.get("merged_columns") or []
        if not merged_cols:
            merged_cols = [
                "date",
                "transaction_value",
                "currency",
                "currency_exchange",
                "transaction_value_idr",
            ]

        for field in merged_cols:
            col_idx = field_to_col.get(field)
            if not col_idx:
                continue

            try:
                ws.merge_cells(
                    start_row=row_start,
                    start_column=col_idx,
                    end_row=row_end,
                    end_column=col_idx,
                )
                _set_merged_top_alignment(ws, row_start, col_idx)
            except Exception:
                pass

def _operation_cost_totals_by_column(expenses, cat_names, category_by_id_map=None):
    totals = [0.0] * len(cat_names)
    for expense in expenses:
        nominal_idr = _expense_amount_for_display(expense)
        col_idx = _map_expense_category_index_from_name(
            _expense_column_mapping_name(expense, category_by_id_map), cat_names
        )
        if col_idx is not None:
            totals[col_idx] += nominal_idr
    return totals

def _get_revenue_bounds(revenues):
    """
    ? CENTRALIZED: Calculate revenue table row bounds.
    Returns (last_revenue_row, total_revenue_row) tuple.

    Handles empty data case by falling back to template row.
    """
    if revenues:
        last_revenue_row = REVENUE_START_ROW + len(revenues) - 1
        total_revenue_row = last_revenue_row + 1
    else:
        # No data: fall back to template row to ensure formulas still work
        last_revenue_row = REVENUE_TEMPLATE_END
        total_revenue_row = REVENUE_TEMPLATE_END
    return last_revenue_row, total_revenue_row

def _group_revenues_by_type(revenues):
    """
    Group revenues by revenue_type (pendapatan_langsung vs pendapatan_lain_lain).
    Returns dict with 'pendapatan_langsung' and 'pendapatan_lain_lain' keys.
    """
    grouped = {"pendapatan_langsung": [], "pendapatan_lain_lain": []}

    for r in revenues:
        r_type = (r.get("revenue_type") or "pendapatan_langsung").strip().lower()
        if r_type == "pendapatan_lain_lain":
            grouped["pendapatan_lain_lain"].append(r)
        else:
            grouped["pendapatan_langsung"].append(r)

    return grouped

def _validate_revenue_data(revenues):
    """? VALIDATION LAYER: Prevent type bugs before render"""
    for r in revenues:
        # CRITICAL: transfer_fee must be numeric, not string like "Pemungut"
        transfer_fee = r.get("transfer_fee")
        if isinstance(transfer_fee, str):
            print(
                f"[VALIDATION ERROR] transfer_fee is string at revenue ID {r.get('id')}: '{transfer_fee}'"
            )
            logger.error(
                f"Validation error: transfer_fee is string at revenue ID {r.get('id')}"
            )
            # Auto-fix: move string to remark, set transfer_fee to 0
            existing_remark = r.get("remark") or ""
            r["remark"] = f"{existing_remark} {transfer_fee}".strip()
            r["transfer_fee"] = 0
        # Ensure numeric types
        r["transfer_fee"] = _to_float(r.get("transfer_fee"), 0)
        r["ppn"] = _to_float(r.get("ppn"), 0)
        r["pph_23"] = _to_float(r.get("pph_23"), 0)
        r["amount_received"] = _to_float(r.get("amount_received"), 0)

