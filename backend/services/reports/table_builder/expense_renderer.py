import logging
logger = logging.getLogger(__name__)
from .expense_renderer_utils import _expense_column_mapping_name, _expense_amount_for_display, _expense_subcategory_label, _group_expenses_by_subcategory
import logging
import re
from datetime import datetime
from io import BytesIO
from typing import Dict, List
from models import (
    Category,
    ManualCombineGroup,
)
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from services.reports.annual_service.annual_constants import (
    COL_AMOUNT_RECEIVED_IDR,
    COL_AMOUNT_RECEIVED_USD,
    COL_CLIENT,
    COL_CURRENCY,
    COL_EXCHANGE_RATE,
    COL_INVOICE_NUMBER,
    COL_INVOICE_VALUE_IDR,
    COL_INVOICE_VALUE_USD,
    COL_PPH_23,
    COL_PPN,
    COL_RECEIVE_DATE,
    COL_REMARK,
    COL_TAX_CURRENCY,
    COL_TAX_DATE,
    COL_TAX_DPP_PEMASUKAN_NEGARA,
    COL_TAX_DPP_PPH21,
    COL_TAX_DPP_PPH23,
    COL_TAX_DPP_PPH26,
    COL_TAX_DPP_PPN,
    COL_TAX_INV_VALUE_USD,
    COL_TAX_RATE,
    COL_TAX_TRANS_VALUE,
    COL_TAX_TRANS_VALUE_IDR,
    COL_TAX_VALUE_PEMASUKAN_NEGARA,
    COL_TAX_VALUE_PPH21,
    COL_TAX_VALUE_PPH23,
    COL_TAX_VALUE_PPH26,
    COL_TAX_VALUE_PPN,
    COL_TRANSFER_FEE,
)
from services.reports.annual_service.annual_helpers import (
    BLUE_FILL,
    GREEN_FILL,
    REVENUE_START_ROW,
    REVENUE_TEMPLATE_END,
    THIN_BORDER,
    _clear_merged_ranges_in_region,
    _root_category_info,
)
from utils.report_helpers import (
    _as_iso_date,
    _clean_currency,
    _clean_settlement_title,
    _clear_range,
    _clear_range_force,
    _extract_imported_row,
    _get_dynamic_row_height,
    _idr_from_currency,
    _is_batch_settlement,
    _map_expense_category_index_from_name,
    _map_expense_column,
    _parse_iso_date,
    _safe_set_cell,
    _safe_set_number,
    _safe_text,
    _set_date_with_format,
    _set_formula_with_format,
    _set_rows_hidden,
    _shorten,
    _to_float,
)





def _render_expense_section_from_data(
    ws,
    expenses: List[Dict],
    cat_names: List[str],
    category_by_id_map: Dict,
    year: int,
    start_row: int = 41,
    sort_mode: str = "kategori",
) -> int:
    from services.reports.annual_service.annual_excel import _clone_row_format

    """
    Render expense section from data.
    """
    logger.debug("Starting data-driven expense rendering")
    last_category_col = 9 + len(cat_names) - 1
    # Define a safe maximum column to clear template "ghost" data (up to column 60/BH)
    SAFE_MAX_COL = 60
    actual_last_col = max(last_category_col, SAFE_MAX_COL)

    no_fill = PatternFill(fill_type=None)
    no_border = Border(
        left=Side(style=None),
        right=Side(style=None),
        top=Side(style=None),
        bottom=Side(style=None),
    )

    white_fill = PatternFill(fill_type="solid", fgColor="FFFFFF")
    green_fill = GREEN_FILL
    blue_fill = BLUE_FILL

    # ? OPTIMIZATION: Cache style objects for the data rows
    # This avoids repeated lookups and copy() calls inside the loop
    data_font = Font(size=11, name="Arial Narrow")
    bold_font = Font(size=11, name="Arial Narrow", bold=True)
    center_align = Alignment(horizontal="center", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # ? STEP 0: Pre-calculate total rows to clear ghost data in BULK
    # This is much faster than clearing row-by-row inside the loop
    # We clear a large block below start_row to ensure a clean slate
    logger.debug("Bulk clearing ghost data area")
    max_cleanup_row = min(start_row + len(expenses) + 200, ws.max_row)
    if max_cleanup_row > start_row:
        _clear_merged_ranges_in_region(
            ws, start_row, max_cleanup_row, 2, actual_last_col
        )
        _clear_range_force(
            ws, start_row, max_cleanup_row, 2, actual_last_col, reset_style=True
        )

    # ? STEP 1: Separate batch vs single expenses

    batch_expenses = [
        e
        for e in expenses
        if _is_batch_settlement(e.get("settlement_type"), e.get("settlement_title"))
    ]
    single_expenses = [
        e
        for e in expenses
        if not _is_batch_settlement(e.get("settlement_type"), e.get("settlement_title"))
    ]

    row_cursor = start_row
    seq_counter = 1

    # ? STEP 2: Render Single Expenses FIRST
    if sort_mode == "tanggal":
        # Sort single expenses by date, then id
        single_expenses.sort(
            key=lambda e: (
                _parse_iso_date(e.get("date"))
                or datetime(year if year > 0 else 2000, 1, 1).date(),
                str(e.get("id") or "0"),
            )
        )
        has_single_data = bool(single_expenses)
    else:
        single_expenses.sort(
            key=lambda e: (
                _extract_imported_row(e.get("notes")) is None,
                _extract_imported_row(e.get("notes")) or 10**9,
                str(e.get("id") or "0"),
            )
        )
        single_grouped = _group_expenses_by_subcategory(single_expenses)
        has_single_data = bool(single_grouped["groups"]) or bool(
            single_grouped["uncategorized"]
        )

    if has_single_data:
        if sort_mode == "tanggal":
            single_start_row = row_cursor
            for idx, expense in enumerate(single_expenses):
                # Optimized: No cloning or clearing per row
                ws.row_dimensions[row_cursor].height = ws.row_dimensions[
                    start_row
                ].height
                for col in range(2, last_category_col + 5):
                    cell = ws.cell(row=row_cursor, column=col)
                    cell.fill = white_fill
                    cell.border = THIN_BORDER
                    cell.font = data_font

                clean_desc = re.sub(
                    r"^\[.*?\]\s*", "", (expense.get("description") or "").strip()
                ).strip()
                _set_date_with_format(ws, row_cursor, 2, expense.get("date"))
                ws.cell(row=row_cursor, column=2).alignment = right_align

                _safe_set_cell(ws, row_cursor, 3, seq_counter)
                ws.cell(row=row_cursor, column=3).alignment = center_align

                _safe_set_cell(ws, row_cursor, 4, clean_desc or "-")
                ws.cell(row=row_cursor, column=4).alignment = left_align

                # ✅ DYNAMIC ROW HEIGHT for Activity (Table 3)
                ws.row_dimensions[row_cursor].height = _get_dynamic_row_height(
                    clean_desc
                )

                _safe_set_cell(ws, row_cursor, 5, expense.get("source") or "-")
                _safe_set_number(ws, row_cursor, 6, _to_float(expense.get("amount")))
                _safe_set_cell(
                    ws, row_cursor, 7, _clean_currency(expense.get("currency"))
                )
                _safe_set_number(
                    ws,
                    row_cursor,
                    8,
                    _to_float(expense.get("currency_exchange"), 1) or 1,
                )
                ws.cell(row=row_cursor, column=7).alignment = center_align
                ws.cell(row=row_cursor, column=8).alignment = center_align

                root_name, _ = _root_category_info(
                    expense.get("category_id"), category_by_id_map
                )
                cat_idx = _map_expense_category_index_from_name(root_name, cat_names)
                if cat_idx is not None:
                    _set_formula_with_format(
                        ws, row_cursor, 9 + cat_idx, f"=$F{row_cursor}*$H{row_cursor}"
                    )

                _set_formula_with_format(
                    ws,
                    row_cursor,
                    last_category_col + 1,
                    f"=$F{row_cursor}*$H{row_cursor}",
                )
                if idx == len(single_expenses) - 1:
                    subtotal_col_letter = get_column_letter(last_category_col + 1)
                    _set_formula_with_format(
                        ws,
                        row_cursor,
                        last_category_col + 2,
                        f"=SUM({subtotal_col_letter}{single_start_row}:{subtotal_col_letter}{row_cursor})",
                    )
                row_cursor += 1
                seq_counter += 1
        else:
            for subcat, items in single_grouped["groups"].items():
                # _clone_row_format replaced by direct assignment for speed
                ws.row_dimensions[row_cursor].height = _get_dynamic_row_height(subcat)
                for col in range(2, last_category_col + 5):
                    cell = ws.cell(row=row_cursor, column=col)
                    cell.fill = white_fill
                    cell.border = THIN_BORDER

                _safe_set_cell(ws, row_cursor, 4, subcat)
                ws.cell(row=row_cursor, column=4).font = bold_font
                ws.cell(row=row_cursor, column=4).alignment = left_align

                row_cursor += 1
                subcat_start_row = row_cursor

                for idx, expense in enumerate(items):
                    # Optimized: No cloning or clearing per row
                    ws.row_dimensions[row_cursor].height = ws.row_dimensions[
                        start_row
                    ].height
                    for col in range(2, last_category_col + 5):
                        cell = ws.cell(row=row_cursor, column=col)
                        cell.fill = white_fill
                        cell.border = THIN_BORDER
                        cell.font = data_font

                    clean_desc = re.sub(
                        r"^\[.*?\]\s*", "", (expense.get("description") or "").strip()
                    ).strip()
                    _set_date_with_format(ws, row_cursor, 2, expense.get("date"))
                    ws.cell(row=row_cursor, column=2).alignment = right_align

                    _safe_set_cell(ws, row_cursor, 3, seq_counter)
                    ws.cell(row=row_cursor, column=3).alignment = center_align

                    _safe_set_cell(ws, row_cursor, 4, clean_desc or "-")
                    ws.cell(row=row_cursor, column=4).alignment = left_align

                    # ✅ DYNAMIC ROW HEIGHT for Activity (Table 3)
                    ws.row_dimensions[row_cursor].height = _get_dynamic_row_height(
                        clean_desc
                    )

                    _safe_set_cell(ws, row_cursor, 5, expense.get("source") or "-")
                    _safe_set_number(
                        ws, row_cursor, 6, _to_float(expense.get("amount"))
                    )
                    _safe_set_cell(
                        ws, row_cursor, 7, _clean_currency(expense.get("currency"))
                    )
                    _safe_set_number(
                        ws,
                        row_cursor,
                        8,
                        _to_float(expense.get("currency_exchange"), 1) or 1,
                    )
                    ws.cell(row=row_cursor, column=7).alignment = center_align
                    ws.cell(row=row_cursor, column=8).alignment = center_align

                    root_name, _ = _root_category_info(
                        expense.get("category_id"), category_by_id_map
                    )
                    cat_idx = _map_expense_category_index_from_name(
                        root_name, cat_names
                    )
                    if cat_idx is not None:
                        _set_formula_with_format(
                            ws,
                            row_cursor,
                            9 + cat_idx,
                            f"=$F{row_cursor}*$H{row_cursor}",
                        )

                    _set_formula_with_format(
                        ws,
                        row_cursor,
                        last_category_col + 1,
                        f"=$F{row_cursor}*$H{row_cursor}",
                    )
                    if idx == len(items) - 1:
                        subtotal_col_letter = get_column_letter(last_category_col + 1)
                        _set_formula_with_format(
                            ws,
                            row_cursor,
                            last_category_col + 2,
                            f"=SUM({subtotal_col_letter}{subcat_start_row}:{subtotal_col_letter}{row_cursor})",
                        )
                    row_cursor += 1
                    seq_counter += 1

            uncat_start_row = row_cursor
            for idx, expense in enumerate(single_grouped["uncategorized"]):
                ws.row_dimensions[row_cursor].height = ws.row_dimensions[
                    start_row
                ].height
                for col in range(2, last_category_col + 5):
                    cell = ws.cell(row=row_cursor, column=col)
                    cell.fill = white_fill
                    cell.border = THIN_BORDER
                    cell.font = data_font

                clean_desc = re.sub(
                    r"^\[.*?\]\s*", "", (expense.get("description") or "").strip()
                ).strip()
                _set_date_with_format(ws, row_cursor, 2, expense.get("date"))
                ws.cell(row=row_cursor, column=2).alignment = right_align

                _safe_set_cell(ws, row_cursor, 3, seq_counter)
                ws.cell(row=row_cursor, column=3).alignment = center_align

                _safe_set_cell(ws, row_cursor, 4, clean_desc or "-")
                ws.cell(row=row_cursor, column=4).alignment = left_align

                # ✅ DYNAMIC ROW HEIGHT for Uncategorized (Table 3)
                ws.row_dimensions[row_cursor].height = _get_dynamic_row_height(
                    clean_desc
                )

                _safe_set_number(ws, row_cursor, 6, _to_float(expense.get("amount")))
                _safe_set_cell(
                    ws, row_cursor, 7, _clean_currency(expense.get("currency"))
                )
                _safe_set_number(
                    ws,
                    row_cursor,
                    8,
                    _to_float(expense.get("currency_exchange"), 1) or 1,
                )
                ws.cell(row=row_cursor, column=7).alignment = center_align
                ws.cell(row=row_cursor, column=8).alignment = center_align

                root_name, _ = _root_category_info(
                    expense.get("category_id"), category_by_id_map
                )
                cat_idx = _map_expense_category_index_from_name(root_name, cat_names)
                if cat_idx is not None:
                    _set_formula_with_format(
                        ws, row_cursor, 9 + cat_idx, f"=$F{row_cursor}*$H{row_cursor}"
                    )

                _set_formula_with_format(
                    ws,
                    row_cursor,
                    last_category_col + 1,
                    f"=$F{row_cursor}*$H{row_cursor}",
                )
                if idx == len(single_grouped["uncategorized"]) - 1:
                    subtotal_col_letter = get_column_letter(last_category_col + 1)
                    _set_formula_with_format(
                        ws,
                        row_cursor,
                        last_category_col + 2,
                        f"=SUM({subtotal_col_letter}{uncat_start_row}:{subtotal_col_letter}{row_cursor})",
                    )

                row_cursor += 1
                seq_counter += 1
    else:
        _clone_row_format(
            ws, start_row, row_cursor, start_col=2, end_col=actual_last_col
        )
        _clear_range(ws, row_cursor, row_cursor, 2, actual_last_col)
        for col in range(2, last_category_col + 5):
            cell = ws.cell(row=row_cursor, column=col)
            cell.border = THIN_BORDER
        _safe_set_cell(ws, row_cursor, 4, "Belum ada data single pengeluaran")
        cell = ws.cell(row=row_cursor, column=4)
        cell.font = cell.font.copy(italic=True, color="808080")
        cell.alignment = cell.alignment.copy(horizontal="center")

        # ? CRITICAL: Reset borders for ALL ghost columns in this row
        for col in range(last_category_col + 5, actual_last_col + 1):
            ws.cell(row=row_cursor, column=col).border = no_border
            ws.cell(row=row_cursor, column=col).fill = no_fill

        row_cursor += 1

    # ? STEP 3: Render separator (green fill)
    separator_row = row_cursor
    _clear_range(ws, separator_row, separator_row, 2, actual_last_col)
    for col in range(2, last_category_col + 5):
        cell = ws.cell(row=separator_row, column=col)
        cell.fill = green_fill
        cell.border = THIN_BORDER

    # Clear ghost borders for separator row
    for col in range(last_category_col + 5, actual_last_col + 1):
        ws.cell(row=separator_row, column=col).border = no_border
        ws.cell(row=separator_row, column=col).fill = no_fill

    _safe_set_cell(ws, separator_row, 4, "OPERATION COST AND OFFICE")
    ws.cell(row=separator_row, column=4).font = Font(bold=True)
    row_cursor += 1

    # Add "Subject :" row
    subject_row = row_cursor
    _clear_range(ws, subject_row, subject_row, 2, actual_last_col)
    for col in range(2, last_category_col + 5):
        cell = ws.cell(row=subject_row, column=col)
        cell.fill = blue_fill
        cell.border = THIN_BORDER
    for col in range(last_category_col + 5, actual_last_col + 1):
        ws.cell(row=subject_row, column=col).border = no_border
        ws.cell(row=subject_row, column=col).fill = no_fill
    _safe_set_cell(ws, subject_row, 2, "Subject")
    ws.cell(row=subject_row, column=2).font = Font(
        bold=True, size=12, name="Arial Narrow"
    )
    _safe_set_cell(ws, subject_row, 3, ":")
    ws.cell(row=subject_row, column=3).font = Font(
        bold=True, size=12, name="Arial Narrow"
    )
    ws.cell(row=subject_row, column=3).alignment = Alignment(
        horizontal="center", vertical="center"
    )
    row_cursor += 1

    # ? STEP 4: Render Batch Expenses SECOND
    batch_by_settlement = {}
    for e in batch_expenses:
        settlement_id = e.get("settlement_id")
        settlement_title = e.get("settlement_title", "Tanpa Settlement")
        key = f"{settlement_id}:{settlement_title}"
        if key not in batch_by_settlement:
            batch_by_settlement[key] = []
        batch_by_settlement[key].append(e)

    sorted_batches = sorted(
        batch_by_settlement.items(),
        key=lambda x: (
            int(x[0].split(":")[0]) if x[0].split(":")[0].isdigit() else 0,
            x[0],
        ),
    )

    if sorted_batches:
        batch_counter = 0
        for settlement_key, batch_items in sorted_batches:
            batch_counter += 1
            settlement_title = _clean_settlement_title(
                batch_items[0].get("settlement_title") or "Tanpa Settlement"
            )

            # Batch header (blue fill)
            batch_header_row = row_cursor
            # No row-by-row clearing here because we did bulk clear at STEP 0

            header_font_large = Font(
                bold=True, size=12, name="Arial Narrow", color="000000"
            )

            for col in range(2, last_category_col + 5):
                cell = ws.cell(row=batch_header_row, column=col)
                cell.fill = blue_fill
                cell.border = THIN_BORDER
                cell.font = header_font_large
                cell.alignment = center_align

            _safe_set_cell(ws, batch_header_row, 2, f"Subject#{batch_counter}")
            ws.cell(row=batch_header_row, column=2).alignment = left_align

            # ? FIX COLON: Bold, Center, Arial Narrow 12
            _safe_set_cell(ws, batch_header_row, 3, ":")
            ws.cell(row=batch_header_row, column=3).alignment = center_align
            ws.cell(row=batch_header_row, column=3).font = header_font_large

            _safe_set_cell(ws, batch_header_row, 4, settlement_title)
            ws.cell(row=batch_header_row, column=4).alignment = left_align
            ws.row_dimensions[batch_header_row].height = max(
                20.0, _get_dynamic_row_height(settlement_title)
            )
            row_cursor += 1

            batch_expenses_start_row = row_cursor

            # Group batch items by subcategory
            batch_grouped = _group_expenses_by_subcategory(batch_items)
            for subcat, items in batch_grouped["groups"].items():
                ws.row_dimensions[row_cursor].height = _get_dynamic_row_height(subcat)
                for col in range(2, last_category_col + 5):
                    cell = ws.cell(row=row_cursor, column=col)
                    cell.fill = white_fill
                    cell.border = THIN_BORDER
                    cell.alignment = center_align

                _safe_set_cell(ws, row_cursor, 4, subcat)
                ws.cell(row=row_cursor, column=4).font = bold_font
                ws.cell(row=row_cursor, column=4).alignment = left_align

                row_cursor += 1
                subcat_start_row = row_cursor

                for idx, expense in enumerate(items):
                    ws.row_dimensions[row_cursor].height = ws.row_dimensions[
                        start_row
                    ].height
                    for col in range(2, last_category_col + 5):
                        cell = ws.cell(row=row_cursor, column=col)
                        cell.fill = white_fill
                        cell.border = THIN_BORDER
                        cell.font = data_font

                    clean_desc = re.sub(
                        r"^\[.*?\]\s*", "", (expense.get("description") or "").strip()
                    ).strip()
                    _set_date_with_format(ws, row_cursor, 2, expense.get("date"))
                    ws.cell(row=row_cursor, column=2).alignment = right_align

                    _safe_set_cell(ws, row_cursor, 3, seq_counter)
                    ws.cell(row=row_cursor, column=3).alignment = center_align

                    _safe_set_cell(ws, row_cursor, 4, clean_desc or "-")
                    ws.cell(row=row_cursor, column=4).alignment = left_align

                    # ✅ DYNAMIC ROW HEIGHT for Batch items (Table 3)
                    ws.row_dimensions[row_cursor].height = _get_dynamic_row_height(
                        clean_desc
                    )

                    _safe_set_cell(ws, row_cursor, 5, expense.get("source") or "-")
                    _safe_set_number(
                        ws, row_cursor, 6, _to_float(expense.get("amount"))
                    )
                    _safe_set_cell(
                        ws, row_cursor, 7, _clean_currency(expense.get("currency"))
                    )
                    _safe_set_number(
                        ws,
                        row_cursor,
                        8,
                        _to_float(expense.get("currency_exchange"), 1) or 1,
                    )
                    ws.cell(row=row_cursor, column=7).alignment = center_align
                    ws.cell(row=row_cursor, column=8).alignment = center_align

                    root_name, _ = _root_category_info(
                        expense.get("category_id"), category_by_id_map
                    )
                    cat_idx = _map_expense_category_index_from_name(
                        root_name, cat_names
                    )
                    if cat_idx is not None:
                        _set_formula_with_format(
                            ws,
                            row_cursor,
                            9 + cat_idx,
                            f"=$F{row_cursor}*$H{row_cursor}",
                        )

                    _set_formula_with_format(
                        ws,
                        row_cursor,
                        last_category_col + 1,
                        f"=$F{row_cursor}*$H{row_cursor}",
                    )

                    row_cursor += 1
                    seq_counter += 1

            if row_cursor > batch_expenses_start_row:
                last_expense_row = row_cursor - 1
                subtotal_col_letter = get_column_letter(last_category_col + 1)
                _set_formula_with_format(
                    ws,
                    last_expense_row,
                    last_category_col + 2,
                    f"=SUM({subtotal_col_letter}{batch_expenses_start_row}:{subtotal_col_letter}{last_expense_row})",
                )
    else:
        _clone_row_format(
            ws, start_row, row_cursor, start_col=2, end_col=actual_last_col
        )
        _clear_range(ws, row_cursor, row_cursor, 2, actual_last_col)
        for col in range(2, last_category_col + 5):
            cell = ws.cell(row=row_cursor, column=col)
            cell.border = THIN_BORDER
        _safe_set_cell(ws, row_cursor, 4, "Belum ada data batch pengeluaran")
        cell = ws.cell(row=row_cursor, column=4)
        cell.font = cell.font.copy(italic=True, color="808080")
        cell.alignment = cell.alignment.copy(horizontal="center")

        # ? CRITICAL: Reset borders for ALL ghost columns in this row
        for col in range(last_category_col + 5, actual_last_col + 1):
            ws.cell(row=row_cursor, column=col).border = no_border
            ws.cell(row=row_cursor, column=col).fill = no_fill

        row_cursor += 1

    # ? STEP 5: Render TOTAL row
    total_row = row_cursor
    _clear_range(ws, total_row, total_row, 2, actual_last_col)
    merge_range = f"B{total_row}:H{total_row}"
    try:
        ws.unmerge_cells(merge_range)
    except Exception:
        pass
    ws.merge_cells(merge_range)
    _safe_set_cell(ws, total_row, 2, "TOTAL")
    ws.cell(row=total_row, column=2).font = Font(bold=True)
    ws.cell(row=total_row, column=4).alignment = Alignment(
        horizontal="right", vertical="center"
    )

    for i in range(len(cat_names) + 2):
        col = 9 + i
        col_letter = get_column_letter(col)
        # ? Use SUM formula instead of static total
        formula = f"=SUM({col_letter}{start_row}:{col_letter}{total_row - 1})"
        _set_formula_with_format(ws, total_row, col, formula)
        ws.cell(row=total_row, column=col).font = Font(bold=True)
        ws.cell(row=total_row, column=col).alignment = Alignment(
            horizontal="right", vertical="center"
        )

    # Clear ghost borders for TOTAL row
    for col in range(last_category_col + 5, actual_last_col + 1):
        ws.cell(row=total_row, column=col).border = no_border
        ws.cell(row=total_row, column=col).fill = no_fill

    for col in range(2, last_category_col + 5):
        cell = ws.cell(row=total_row, column=col)
        cell.border = THIN_BORDER

    for row in range(start_row, total_row + 1):
        for col in range(2, last_category_col + 5):
            cell = ws.cell(row=row, column=col)
            cell.border = THIN_BORDER

    _set_rows_hidden(ws, start_row, total_row, False)
    max_row = ws.max_row
    if total_row < max_row:
        _set_rows_hidden(ws, total_row + 1, max_row, True)

    logger.debug("Expense rendering completed successfully")
    return total_row

