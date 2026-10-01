import io
import os
import logging
from io import BytesIO
import xlsxwriter
from datetime import datetime

logger = logging.getLogger(__name__)
from config import BASE_DIR
from fastapi.responses import StreamingResponse
from models import Advance, AdvanceItem, Settlement, Expense, Category, db
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as OpenpyxlImage
from utils.parsing_helper import _default_report_year, _parse_iso_date
from .data_helpers import _validate_date, _get_summary_approved_expenses, _build_summary_payload


def generate_excel_report(company, year, mode, start_date, end_date):
    company = company or "EWI Foundation"

    # âœ… Update call dengan parameter baru dan perbaikan jumlah argumen
    expenses = _get_summary_approved_expenses(
        start_date=start_date, end_date=end_date, year=year, mode=mode, company=company
    )
    summary_data, grand_total_val = _build_summary_payload(expenses, company=company)

    wb = Workbook()
    ws = wb.active
    ws.title = "Laporan Summary Bulanan"

    # --- HEADER SECTION ---
    # Row 1: Title
    ws["A1"] = "Daily Report - Balance Calculation Operation Cost"
    ws["A1"].font = Font(bold=True, size=14)

    # Row 2: Project
    ws["A2"] = "Project"
    ws["B2"] = ":"
    ws["C2"] = f"REVENUE vs OPERATION COST Tahun {year}"
    ws["A2"].font = Font(bold=True)
    ws["C2"].font = Font(bold=True)

    # Row 3: Date
    ws["A3"] = "Date"
    ws["B3"] = ":"
    if start_date and end_date:
        ws["C3"] = f"{start_date.strftime('%d %B')} - {end_date.strftime('%d %B %Y')}"
    else:
        ws["C3"] = f"Januari - Desember {year}"
    ws["A3"].font = Font(bold=True)
    ws["C3"].font = Font(bold=True)

    # --- LOGO SECTION ---
    logo_path = os.path.abspath(
        os.path.join(
            BASE_DIR, "..", "frontend", "assets", "images", "sheet1.png"
        )
    )
    if os.path.exists(logo_path):
        img = OpenpyxlImage(logo_path)
        # Sesuai analisis screenshot: Width 6.69cm (~253px), Height 2.22cm (~84px)
        img.width = 253
        img.height = 84
        ws.add_image(img, "N1")  # Anchor di N1 agar sejajar ke arah kolom O

    # --- TABLE SECTION ---
    months = [
        "Jan",
        "Feb",
        "Mar",
        "Apr",
        "Mei",
        "Jun",
        "Jul",
        "Agt",
        "Sep",
        "Okt",
        "Nov",
        "Des",
    ]
    headers = ["Kategori"] + months + ["Total"]

    start_row = 6
    for i, header in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=i)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            start_color="3B82F6", end_color="3B82F6", fill_type="solid"
        )  # Blue header
        cell.alignment = Alignment(horizontal="center")

    current_row = start_row + 1
    for row in summary_data:
        is_parent = row.get("is_parent", False)
        level = row.get("level", 0)
        cat_name = row.get("category", "-")

        display_name = ("  " * level) + cat_name
        ws.cell(row=current_row, column=1).value = display_name

        if is_parent:
            ws.cell(row=current_row, column=1).font = Font(bold=True)

        monthly = row.get("monthly", {})
        for i in range(1, 13):
            val = monthly.get(str(i), 0)
            cell = ws.cell(row=current_row, column=i + 1)
            cell.value = val
            cell.number_format = "#,##0"
            if is_parent:
                cell.font = Font(bold=True)

        total_val = row.get("yearly_total", 0)
        total_cell = ws.cell(row=current_row, column=14)
        total_cell.value = total_val
        total_cell.number_format = "#,##0"
        if is_parent:
            total_cell.font = Font(bold=True)

        current_row += 1

    # Grand Total row
    ws.cell(row=current_row, column=1).value = "GRAND TOTAL"
    ws.cell(row=current_row, column=1).font = Font(bold=True)

    for i in range(1, 13):
        total_m = sum(
            r["monthly"].get(str(i), 0) for r in summary_data if r.get("level") == 0
        )  # Only sum top level to avoid double count
        cell = ws.cell(row=current_row, column=i + 1)
        cell.value = total_m
        cell.number_format = "#,##0"
        cell.font = Font(bold=True)
        cell.fill = PatternFill(
            start_color="E6F0FF", end_color="E6F0FF", fill_type="solid"
        )

    gt_cell = ws.cell(row=current_row, column=14)
    gt_cell.value = grand_total_val
    gt_cell.number_format = "#,##0"
    gt_cell.font = Font(bold=True)
    gt_cell.fill = PatternFill(
        start_color="E6F0FF", end_color="E6F0FF", fill_type="solid"
    )

    # Column widths
    ws.column_dimensions["A"].width = 40
    for i in range(2, 14):
        ws.column_dimensions[get_column_letter(i)].width = 15
    ws.column_dimensions["N"].width = 18

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="summary_{timestamp}.xlsx"'},
    )

def generate_excel_advance_report(advance_id, start_date_str, end_date_str):
    # Natural Sort: panjang kode -> isi kode (A1, A2... A10)
    # Sertakan semua kategori agar data di kategori induk ikut muncul
    categories = Category.query.order_by(
        db.func.length(Category.code).asc(), Category.code.asc()
    ).all()

    category_map = {cat.id: cat for cat in categories}

    query = (
        Advance.query.filter_by(id=advance_id)
        if advance_id
        else Advance.query.filter_by(status="completed")
    )

    if start_date_str:
        try:
            query = query.filter(
                Advance.approved_at >= datetime.strptime(start_date_str, "%Y-%m-%d")
            )
        except ValueError as e:
            logger.debug(f"Invalid start_date format: {start_date_str} - {e}")

    if end_date_str:
        try:
            end_date = datetime.strptime(end_date_str, "%Y-%m-%d").replace(
                hour=23, minute=59, second=59
            )
            query = query.filter(Advance.approved_at <= end_date)
        except ValueError as e:
            logger.debug(f"Invalid end_date format: {end_date_str} - {e}")

    approved_advances = query.order_by(Advance.approved_at.asc()).all()

    # Hitung total nominal per kategori untuk penyembunyian dinamis (Smart Excel)
    category_totals = {}
    for advance in approved_advances:
        for item in advance.items:
            if item.category_id in category_map:
                cn = category_map[item.category_id].name
                category_totals[cn] = category_totals.get(cn, 0) + (item.idr_amount or 0)
                
    # Hanya sertakan kategori yang memiliki transaksi > 0 (Hide Empty Columns)
    filtered_categories = [cat for cat in categories if category_totals.get(cat.name, 0) > 0]
    
    # Jika tidak ada kategori yang aktif sama sekali, sediakan minimal 1 slot reserve
    if not filtered_categories and categories:
        filtered_categories = categories[:3]  # Fallback ke 3 kategori pertama jika semuanya kosong
        
    cat_col_map = {cat.name: i + 4 for i, cat in enumerate(filtered_categories)}

    wb = Workbook()
    ws = wb.active
    ws.title = "Laporan Perencanaan Kasbon"
    headers = ["Tanggal Disetujui", "Deskripsi Kasbon", "Karyawan", "Judul Kasbon"] + [
        cat.name for cat in filtered_categories
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for advance in approved_advances:
        row_data = [None] * len(headers)
        row_data[0] = advance.approved_at if advance.approved_at else None
        row_data[1] = advance.description
        row_data[2] = advance.requester.full_name if advance.requester else ""
        row_data[3] = advance.title
        cat_sums = {}
        for item in advance.items:
            cn = (
                category_map[item.category_id].name
                if item.category_id in category_map
                else None
            )
            if cn:
                cat_sums[cn] = cat_sums.get(cn, 0) + (item.idr_amount or 0)
        for cn, amt in cat_sums.items():
            if cn in cat_col_map:
                row_data[cat_col_map[cn]] = amt
        ws.append(row_data)
    for i, col_name in enumerate(headers, 1):
        w = len(col_name) + 4
        if w > 35:
            w = 35
        ws.column_dimensions[get_column_letter(i)].width = w
    for col, w in {"A": 18, "B": 35, "C": 20, "D": 30}.items():
        ws.column_dimensions[col].width = w
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="laporan_kasbon_{timestamp}.xlsx"'},
    )

