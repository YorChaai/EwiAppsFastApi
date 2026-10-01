import openpyxl, datetime as dt, psycopg2
from psycopg2.extras import DictCursor

EXCEL = r'data/excel/20250427_EWI Financial-Repport_2024-2026 (version 2).xlsx'
wb_f = openpyxl.load_workbook(EXCEL, data_only=False)
wb_v = openpyxl.load_workbook(EXCEL, data_only=True)

def to_num(val):
    if val is None: return 0.0
    if isinstance(val, (int, float)): return float(val)
    s = str(val).strip().replace(',', '')
    try: return float(s)
    except: return 0.0

def _is_marker_row(val_b, val_c):
    b = str(val_b or '').strip().lower()
    c = str(val_c or '').strip()
    return b in ('single', 'batch', 'subject', 'detail') and c == ':'

def _is_subkategori_row(val_b, val_c, val_d, val_e, val_f, val_g):
    d_filled = bool(val_d and str(val_d).strip())
    others_empty = not any([
        val_b and str(val_b).strip(),
        val_c and str(val_c).strip(),
        val_e and str(val_e).strip(),
        val_f and str(val_f).strip(),
        val_g and str(val_g).strip(),
    ])
    return d_filled and others_empty

def parse_sheet_exact(ws_f, ws_v, year):
    mode = None  # Hanya aktif setelah marker Single: atau Batch:
    current_subject = None
    batch_groups = []
    single_items = []
    batch_items = []

    for r in range(1, ws_v.max_row + 1):
        val_b = ws_v.cell(r, 2).value
        val_c = ws_v.cell(r, 3).value
        val_d = ws_v.cell(r, 4).value
        val_e = ws_v.cell(r, 5).value
        val_f = ws_v.cell(r, 6).value
        val_g = ws_v.cell(r, 7).value
        val_h = ws_v.cell(r, 8).value

        val_b_str = str(val_b or '').strip()
        val_c_str = str(val_c or '').strip()
        val_d_str = str(val_d or '').strip()

        all_empty = not any([str(x or '').strip() for x in [val_b, val_c, val_d, val_e, val_f, val_g]])
        if all_empty:
            continue

        if (
            val_b_str.upper() in ('DATE', 'PENGELUARAN')
            or val_c_str.upper() in ('#',)
            or val_d_str.upper() in ('ACTIVITY', 'DESCRIPTION', 'DETAIL', 'ITEM', 'JUMLAH', 'TOTAL', 'SUBTOTAL', 'USD', 'IDR')
        ):
            continue

        if _is_marker_row(val_b, val_c):
            m = val_b_str.lower()
            if m == 'single':
                mode = 'single'
                current_subject = None
            elif m == 'batch':
                mode = 'batch'
                current_subject = None
            elif m == 'subject':
                if mode == 'batch':
                    current_subject = val_d_str[:100]
                    batch_groups.append(current_subject)
            elif m == 'detail':
                pass
            continue

        if 'TRANSFER FROM' in val_b_str.upper():
            continue

        if _is_subkategori_row(val_b, val_c, val_d, val_e, val_f, val_g):
            continue

        if mode is None:
            continue

        # Amount extraction exactly like V14
        col_amt = 0.0
        found_in_col = False
        for col in range(9, 22):
            v_tmp = to_num(ws_v.cell(r, col).value)
            if v_tmp != 0:
                col_amt = v_tmp
                found_in_col = True
                break

        if not found_in_col:
            for fallback_col in [6, 5]:
                v_tmp = to_num(ws_v.cell(r, fallback_col).value)
                if v_tmp != 0:
                    col_amt = v_tmp
                    found_in_col = True
                    break

        has_date = isinstance(val_b, (dt.date, dt.datetime))
        cur_raw = str(val_g or '').strip().upper()
        if cur_raw in ('RP', ''):
            cur = 'IDR'
        else:
            cur = cur_raw

        if not found_in_col and has_date and cur in ('IDR', 'USD', 'MYR'):
            col_amt = 0.0
            found_in_col = True

        if not found_in_col:
            continue

        if val_d_str.upper() in ('TOTAL', 'SUBTOTAL', '-', ''):
            continue

        rat_val = to_num(ws_v.cell(r, 8).value)
        calc_rat = rat_val if rat_val > 0 else 1.0
        if cur in ('IDR', '-', ''):
            calc_rat = 1.0

        if cur not in ('IDR', '-', '') and calc_rat > 1:
            t_st = round(col_amt / calc_rat, 4)
        else:
            t_st = col_amt

        item_amt_idr = t_st * calc_rat

        if mode == 'batch':
            batch_items.append((r, item_amt_idr))
        else:
            single_items.append((r, item_amt_idr))

    return {
        'single_count': len(single_items),
        'single_amt': sum(x[1] for x in single_items),
        'batch_count': len(batch_groups),
        'batch_items_count': len(batch_items),
        'batch_amt': sum(x[1] for x in batch_items),
    }

conn = psycopg2.connect('postgresql://postgres:yorchai12@localhost:5432/miniproject_revisi_db')
cur = conn.cursor(cursor_factory=DictCursor)

for yr, sname in [(2024, 'Rev VS Ops cost th 2024'), (2025, 'Rev VS Ops cost th 2025_Final'), (2026, 'Rev VS Ops cost th 2026_')]:
    stats = parse_sheet_exact(wb_f[sname], wb_v[sname], yr)
    cur.execute("SELECT count(*), coalesce(sum(disbursed_amount),0) FROM settlements WHERE report_year=%s AND settlement_type='single'", (yr,))
    db_s_cnt, db_s_amt = cur.fetchone()
    cur.execute("SELECT count(*), coalesce(sum(disbursed_amount),0) FROM settlements WHERE report_year=%s AND settlement_type='batch'", (yr,))
    db_b_cnt, db_b_amt = cur.fetchone()
    cur.execute("SELECT count(*) FROM expenses e JOIN settlements s ON e.settlement_id=s.id WHERE s.report_year=%s AND s.settlement_type='batch'", (yr,))
    db_b_items = cur.fetchone()[0]

    print(f"=== TAHUN {yr} ===")
    print(f"Single Items: Excel={stats['single_count']} vs DB={db_s_cnt} | Selisih = {stats['single_count'] - db_s_cnt}")
    print(f"Batch Sets  : Excel={stats['batch_count']} vs DB={db_b_cnt} | Selisih = {stats['batch_count'] - db_b_cnt}")
    print(f"Batch Items : Excel={stats['batch_items_count']} vs DB={db_b_items} | Selisih = {stats['batch_items_count'] - db_b_items}")
    print(f"Single Amt  : Excel=Rp {stats['single_amt']:,.2f} vs DB=Rp {float(db_s_amt):,.2f} | Selisih = Rp {stats['single_amt'] - float(db_s_amt):,.2f}")
    print(f"Batch Amt   : Excel=Rp {stats['batch_amt']:,.2f} vs DB=Rp {float(db_b_amt):,.2f} | Selisih = Rp {stats['batch_amt'] - float(db_b_amt):,.2f}")
    print()
