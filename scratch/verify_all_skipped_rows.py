import openpyxl
import datetime as dt

excel_path = r"D:\Aplikasi\EWIApps\data\excel\20250427_EWI Financial-Repport_2024-2026 (version 2).xlsx"
sheet_name = "Rev VS Ops cost th 2026_"

wb_val = openpyxl.load_workbook(excel_path, data_only=True)
wb_form = openpyxl.load_workbook(excel_path, data_only=False)

ws_v = wb_val[sheet_name]
ws_f = wb_form[sheet_name]

from backend.scripts.generate_hybrid_migration_v14 import to_num_cell

skipped_batch_rows = []
valid_batch_rows = []

for r in range(391, 831):
    val_a = ws_v.cell(r, 1).value
    val_b = ws_f.cell(r, 2).value
    val_c = ws_f.cell(r, 3).value
    val_d = ws_f.cell(r, 4).value

    # Check if there is an amount
    col_amt = 0
    found_in_col = False
    for col in range(5, 30):
        v_tmp = to_num_cell(ws_f, ws_v, r, col)
        if v_tmp != 0:
            col_amt = v_tmp
            found_in_col = True
            break

    desc_for_check = str(val_d or val_c or "").strip().upper()
    if desc_for_check in ["TOTAL", "SUBTOTAL", "-", ""]:
        skipped_batch_rows.append((r, "TOTAL/SUBTOTAL", desc_for_check))
        continue

    # Header noises
    val_d_str = str(val_d or "").upper().strip()
    val_c_str = str(val_c or "").upper().strip()
    val_b_str = str(val_b or "").upper().strip()
    val_a_str = str(val_a or "").upper().strip()

    if any(h in val_d_str or h in val_c_str for h in ["DETAIL/DESCRIPTION", "ACTIVITY", "DESCRIPTION"]) or any(h in val_b_str or h in val_a_str for h in ["DATE", "#"]):
        skipped_batch_rows.append((r, "TABLE HEADER NOISE", f"b={val_b} d={val_d}"))
        continue

    if found_in_col:
        valid_batch_rows.append((r, val_d or val_c, col_amt))
    else:
        # Check why it's skipped
        non_empty_cells = [(c, ws_v.cell(r, c).value) for c in range(1, 15) if ws_v.cell(r, c).value is not None]
        if not non_empty_cells:
            skipped_batch_rows.append((r, "EMPTY ROW", ""))
        else:
            skipped_batch_rows.append((r, "TITLE/LABEL/OTHER", str(non_empty_cells)[:60]))

print(f"Total Baris 391-830: {830 - 391 + 1} baris")
print(f"Valid Batch Items: {len(valid_batch_rows)} baris")
print(f"Skipped Rows in Batch Section: {len(skipped_batch_rows)} baris")

print("\nRincian Baris yang Di-skip dalam Bagian Batch (391-830):")
for r, reason, text in skipped_batch_rows:
    print(f"  Row {r:3d}: [{reason:20s}] {text}")
