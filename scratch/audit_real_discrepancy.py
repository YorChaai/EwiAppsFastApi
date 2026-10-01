import openpyxl
import psycopg2
from psycopg2.extras import DictCursor

wb_val = openpyxl.load_workbook('data/excel/20250427_EWI Financial-Repport_2024-2026 (version 2).xlsx', data_only=True)
conn = psycopg2.connect('postgresql://postgres:yorchai12@localhost:5432/miniproject_revisi_db')
cur = conn.cursor(cursor_factory=DictCursor)

for yr in [2024, 2025, 2026]:
    sname = {
        2024: 'Rev VS Ops cost th 2024',
        2025: 'Rev VS Ops cost th 2025_Final',
        2026: 'Rev VS Ops cost th 2026_'
    }[yr]
    ws = wb_val[sname]
    
    col_range = (8, 21) if yr == 2024 else (9, 22)
    start_r = 71 if yr == 2024 else (162 if yr == 2025 else 176)
    end_r = 861 if yr == 2024 else (1305 if yr == 2025 else 830)
    
    excel_items = []
    for r in range(start_r, end_r + 1):
        desc = str(ws.cell(r, 4).value or ws.cell(r, 3).value or '').strip()
        if desc.upper() in ['TOTAL', 'SUBTOTAL', 'SUB TOTAL', '-']:
            continue
        s = sum([ws.cell(r, c).value for c in range(col_range[0], col_range[1]) if isinstance(ws.cell(r, c).value, (int, float))])
        if s > 0:
            excel_items.append((r, desc, s))
            
    cur.execute('''
        SELECT e.id, e.description, e.amount, e.currency, e.currency_exchange, 
               (e.amount * e.currency_exchange) as idr_amt, s.settlement_type, s.report_year
        FROM expenses e
        JOIN settlements s ON e.settlement_id = s.id
        WHERE s.report_year = %s
        ORDER BY e.id ASC;
    ''', (yr,))
    db_items = cur.fetchall()
    
    excel_tot = sum([x[2] for x in excel_items])
    db_tot = sum([float(x['idr_amt']) for x in db_items])
    
    print(f"\n{'='*70}")
    print(f"   AUDIT TAHUN {yr} (Sheet: {sname})")
    print(f"{'='*70}")
    print(f"Excel Items: {len(excel_items):4d} items | Total IDR: Rp {excel_tot:,.2f}")
    print(f"DB Items:    {len(db_items):4d} items | Total IDR: Rp {db_tot:,.2f}")
    print(f"SELISIH (Excel - DB): Rp {excel_tot - db_tot:,.2f}")
    
    # Check item count diff
    matched_excel = set()
    matched_db = set()
    
    for idx_e, (r, desc_e, amt_e) in enumerate(excel_items):
        for idx_d, d in enumerate(db_items):
            if idx_d in matched_db:
                continue
            if abs(float(d['idr_amt']) - float(amt_e)) < 0.05:
                matched_excel.add(idx_e)
                matched_db.add(idx_d)
                break
                
    unmatched_excel = [excel_items[i] for i in range(len(excel_items)) if i not in matched_excel]
    unmatched_db = [db_items[i] for i in range(len(db_items)) if i not in matched_db]
    
    print(f"Matched count: {len(matched_excel)} items")
    print(f"Unmatched in Excel: {len(unmatched_excel)} items (Total: Rp {sum(x[2] for x in unmatched_excel):,.2f})")
    print(f"Unmatched in DB:    {len(unmatched_db)} items (Total: Rp {sum(float(x['idr_amt']) for x in unmatched_db):,.2f})")
    
    if unmatched_excel:
        print("  Sampel Unmatched Excel (5 pertama):")
        for r, desc, amt in unmatched_excel[:5]:
            print(f"    [Row {r}] {desc[:40]}: Rp {amt:,.2f}")
    if unmatched_db:
        print("  Sampel Unmatched DB (5 pertama):")
        for d in unmatched_db[:5]:
            print(f"    [ID {d['id']}] {d['description'][:40]}: raw={d['amount']} {d['currency']} @ {d['currency_exchange']} = Rp {d['idr_amt']:,.2f}")

conn.close()
