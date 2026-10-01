import openpyxl
import datetime as dt

excel_path = r"D:\Aplikasi\EWIApps\data\excel\20250427_EWI Financial-Repport_2024-2026 (version 2).xlsx"

wb_val = openpyxl.load_workbook(excel_path, data_only=True)
wb_form = openpyxl.load_workbook(excel_path, data_only=False)

def audit_sheet(sheet_name):
    print(f"\n=======================================================")
    print(f"             AUDIT SHEET: {sheet_name}")
    print(f"=======================================================")
    ws_v = wb_val[sheet_name]
    ws_f = wb_form[sheet_name]
    print(f"Max row: {ws_v.max_row}, Max col: {ws_v.max_column}")

    # Find PENGELUARAN
    r_pengeluaran = -1
    for r in range(1, min(ws_v.max_row + 1, 300)):
        row_txt = " ".join([str(ws_v.cell(r, c).value or "") for c in range(1, 6)]).upper()
        if "PENGELUARAN" in row_txt:
            r_pengeluaran = r
            print(f"Found 'PENGELUARAN' at row {r}: {row_txt[:60]}")
            break

    # Find divider to Batch
    r_divider = -1
    for r in range(r_pengeluaran + 1, min(ws_v.max_row + 1, 1000)):
        row_txt = " ".join([str(ws_v.cell(r, c).value or "") for c in range(1, 6)]).upper()
        if any(kw in row_txt for kw in ["OPERATION COST AND OFFICE", "EXPENSES REPORT", "SUBJECT :", "SUBJECT:"]):
            r_divider = r
            print(f"Found Batch start / divider at row {r}: {row_txt[:60]}")
            break

    print(f"Single Mode Range estimated: Row {r_pengeluaran + 1} to Row {r_divider - 1}")
    print(f"Batch Mode Range estimated : Row {r_divider} to Row {ws_v.max_row}")

audit_sheet("Rev VS Ops cost th 2024")
audit_sheet("Rev VS Ops cost th 2025_Final")
