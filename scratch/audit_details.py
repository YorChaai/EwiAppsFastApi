import openpyxl
import psycopg2
from psycopg2.extras import DictCursor
import datetime as dt

wb_val = openpyxl.load_workbook('data/excel/20250427_EWI Financial-Repport_2024-2026 (version 2).xlsx', data_only=True)
conn = psycopg2.connect('postgresql://postgres:yorchai12@localhost:5432/miniproject_revisi_db')
cur = conn.cursor(cursor_factory=DictCursor)

def audit_year_2024():
    ws = wb_val['Rev VS Ops cost th 2024']
    cur.execute('''
        SELECT e.id, s.title as set_title, s.description as set_desc, s.settlement_type, e.description, 
               e.amount, e.currency, e.currency_exchange, (e.amount * e.currency_exchange) as idr_amt
        FROM expenses e
        JOIN settlements s ON e.settlement_id = s.id
        WHERE s.report_year = 2024
        ORDER BY e.id ASC;
    ''')
    db_items = cur.fetchall()
    db_single = [x for x in db_items if x['settlement_type'] == 'single']
    db_batch = [x for x in db_items if x['settlement_type'] == 'batch']
    
    # Excel Single: rows 71 to 263
    single_excel = []
    for r in range(71, 264):
        desc = str(ws.cell(r, 4).value or '').strip()
        s = sum([ws.cell(r, c).value for c in range(8, 21) if isinstance(ws.cell(r, c).value, (int, float))])
        if s > 0:
            single_excel.append((r, desc, s))
            
    print(f"=== 2024 SINGLE ===")
    print(f"Excel Single: count={len(single_excel)}, sum=Rp {sum(x[2] for x in single_excel):,.2f}")
    print(f"DB Single:    count={len(db_single)}, sum=Rp {sum(float(x['idr_amt']) for x in db_single):,.2f}")
    print(f"Diff: Rp {sum(x[2] for x in single_excel) - sum(float(x['idr_amt']) for x in db_single):,.2f}")
    
    # Excel Batch: rows 264 to 861
    batch_excel = []
    tot_excel_batch_idr = 0.0
    for r in range(264, 862):
        desc = str(ws.cell(r, 4).value or ws.cell(r, 3).value or '').strip()
        cur_val = str(ws.cell(r, 6).value or 'IDR').strip().upper()
        rat_val = ws.cell(r, 7).value
        rate = float(rat_val) if isinstance(rat_val, (int, float)) and rat_val > 0 else 1.0
        if cur_val in ['IDR', '-', '']:
            rate = 1.0
            
        s = sum([ws.cell(r, c).value for c in range(8, 21) if isinstance(ws.cell(r, c).value, (int, float))])
        amt_col5 = ws.cell(r, 5).value if isinstance(ws.cell(r, 5).value, (int, float)) else 0
        
        # If s == 0, check col 5
        amt = s if s != 0 else amt_col5
        if amt != 0:
            idr = amt * rate
            tot_excel_batch_idr += idr
            batch_excel.append((r, desc, amt, cur_val, rate, idr))
            
    # Match exactly
    excel_batch = []
    for r in range(265, 862):
        s = sum([ws.cell(r, c).value for c in range(8, 21) if isinstance(ws.cell(r, c).value, (int, float))])
        desc = str(ws.cell(r, 4).value or '').strip()
        if s != 0:
            excel_batch.append((r, desc, s))

    matched_db = set()
    unmatched_excel = []
    for r, desc, amt in excel_batch:
        found = False
        for idx, d in enumerate(db_batch):
            if idx in matched_db:
                continue
            if abs(float(d['idr_amt']) - float(amt)) < 0.01:
                matched_db.add(idx)
                found = True
                break
        if not found:
            unmatched_excel.append((r, desc, amt))

    print(f"\nUnmatched Excel Batch items: {len(unmatched_excel)}")
    for r, desc, amt in unmatched_excel:
        print(f"  Row {r}: {desc} | Rp {amt:,.2f}")

    unmatched_db = [db_batch[i] for i in range(len(db_batch)) if i not in matched_db]
    print(f"Unmatched DB Batch items: {len(unmatched_db)}")
    for d in unmatched_db:
        print(f"  ID {d['id']}: {d['description']} | Rp {d['idr_amt']:,.2f}")

def audit_year_2025():
    ws = wb_val['Rev VS Ops cost th 2025_Final']
    cur.execute('''
        SELECT e.id, s.title as set_title, s.description as set_desc, s.settlement_type, e.description, 
               e.amount, e.currency, e.currency_exchange, (e.amount * e.currency_exchange) as idr_amt
        FROM expenses e
        JOIN settlements s ON e.settlement_id = s.id
        WHERE s.report_year = 2025
        ORDER BY e.id ASC;
    ''')
    db_items = cur.fetchall()
    db_single = [x for x in db_items if x['settlement_type'] == 'single']
    db_batch = [x for x in db_items if x['settlement_type'] == 'batch']

    # Excel Single: rows 162 to 516
    single_excel = []
    for r in range(162, 517):
        desc = str(ws.cell(r, 4).value or '').strip()
        cur_val = str(ws.cell(r, 7).value or 'IDR').strip().upper()
        rate_val = ws.cell(r, 8).value
        rate = float(rate_val) if isinstance(rate_val, (int, float)) and rate_val > 0 else 1.0
        if cur_val in ['IDR', 'RP', '-', '']:
            rate = 1.0
            
        s = sum([ws.cell(r, c).value for c in range(9, 22) if isinstance(ws.cell(r, c).value, (int, float))])
        amt_col6 = ws.cell(r, 6).value if isinstance(ws.cell(r, 6).value, (int, float)) else 0
        amt = s if s != 0 else amt_col6
        if amt != 0:
            single_excel.append((r, desc, amt * rate))

    print(f"\n{'='*60}\n=== 2025 SINGLE ===")
    print(f"Excel Single: count={len(single_excel)}, sum IDR=Rp {sum(x[2] for x in single_excel):,.2f}")
    print(f"DB Single:    count={len(db_single)}, sum IDR=Rp {sum(float(x['idr_amt']) for x in db_single):,.2f}")
    print(f"Diff: Rp {sum(x[2] for x in single_excel) - sum(float(x['idr_amt']) for x in db_single):,.2f}")

    # Excel Batch: rows 517 to 1307
    batch_excel = []
    for r in range(517, 1308):
        desc = str(ws.cell(r, 4).value or ws.cell(r, 3).value or '').strip()
        if desc.upper() in ['TOTAL', 'SUBTOTAL', 'SUB TOTAL', '-']:
            continue
        cur_val = str(ws.cell(r, 7).value or 'IDR').strip().upper()
        rate_val = ws.cell(r, 8).value
        rate = float(rate_val) if isinstance(rate_val, (int, float)) and rate_val > 0 else 1.0
        if cur_val in ['IDR', 'RP', '-', '']:
            rate = 1.0
        s = sum([ws.cell(r, c).value for c in range(9, 22) if isinstance(ws.cell(r, c).value, (int, float))])
        amt_col6 = ws.cell(r, 6).value if isinstance(ws.cell(r, 6).value, (int, float)) else 0
        amt = s if s != 0 else amt_col6
        if amt != 0:
            batch_excel.append((r, desc, amt * rate))

    print(f"\n=== 2025 BATCH ===")
    print(f"Excel Batch: count={len(batch_excel)}, sum IDR=Rp {sum(x[2] for x in batch_excel):,.2f}")
    print(f"DB Batch:    count={len(db_batch)}, sum IDR=Rp {sum(float(x['idr_amt']) for x in db_batch):,.2f}")
    print(f"Diff: Rp {sum(x[2] for x in batch_excel) - sum(float(x['idr_amt']) for x in db_batch):,.2f}")

def audit_year_2026():
    ws = wb_val['Rev VS Ops cost th 2026_']
    cur.execute('''
        SELECT e.id, s.title as set_title, s.description as set_desc, s.settlement_type, e.description, 
               e.amount, e.currency, e.currency_exchange, (e.amount * e.currency_exchange) as idr_amt
        FROM expenses e
        JOIN settlements s ON e.settlement_id = s.id
        WHERE s.report_year = 2026
        ORDER BY e.id ASC;
    ''')
    db_items = cur.fetchall()
    db_single = [x for x in db_items if x['settlement_type'] == 'single']
    db_batch = [x for x in db_items if x['settlement_type'] == 'batch']

    # Excel Single: rows 176 to 390
    single_excel = []
    for r in range(176, 391):
        desc = str(ws.cell(r, 4).value or '').strip()
        cur_val = str(ws.cell(r, 7).value or 'IDR').strip().upper()
        rate_val = ws.cell(r, 8).value
        rate = float(rate_val) if isinstance(rate_val, (int, float)) and rate_val > 0 else 1.0
        if cur_val in ['IDR', 'RP', '-', '']:
            rate = 1.0
            
        s = sum([ws.cell(r, c).value for c in range(9, 22) if isinstance(ws.cell(r, c).value, (int, float))])
        amt_col6 = ws.cell(r, 6).value if isinstance(ws.cell(r, 6).value, (int, float)) else 0
        amt = s if s != 0 else amt_col6
        if amt != 0:
            single_excel.append((r, desc, amt * rate))

    print(f"\n{'='*60}\n=== 2026 SINGLE ===")
    print(f"Excel Single: count={len(single_excel)}, sum IDR=Rp {sum(x[2] for x in single_excel):,.2f}")
    print(f"DB Single:    count={len(db_single)}, sum IDR=Rp {sum(float(x['idr_amt']) for x in db_single):,.2f}")
    print(f"Diff: Rp {sum(x[2] for x in single_excel) - sum(float(x['idr_amt']) for x in db_single):,.2f}")

    # Excel Batch: rows 391 to 830
    batch_excel = []
    for r in range(391, 831):
        desc = str(ws.cell(r, 4).value or ws.cell(r, 3).value or '').strip()
        if desc.upper() in ['TOTAL', 'SUBTOTAL', 'SUB TOTAL', '-']:
            continue
        cur_val = str(ws.cell(r, 7).value or 'IDR').strip().upper()
        rate_val = ws.cell(r, 8).value
        rate = float(rate_val) if isinstance(rate_val, (int, float)) and rate_val > 0 else 1.0
        if cur_val in ['IDR', 'RP', '-', '']:
            rate = 1.0
        s = sum([ws.cell(r, c).value for c in range(9, 22) if isinstance(ws.cell(r, c).value, (int, float))])
        amt_col6 = ws.cell(r, 6).value if isinstance(ws.cell(r, 6).value, (int, float)) else 0
        amt = s if s != 0 else amt_col6
        if amt != 0:
            batch_excel.append((r, desc, amt * rate))

    print(f"\n=== 2026 BATCH ===")
    print(f"Excel Batch: count={len(batch_excel)}, sum IDR=Rp {sum(x[2] for x in batch_excel):,.2f}")
    print(f"DB Batch:    count={len(db_batch)}, sum IDR=Rp {sum(float(x['idr_amt']) for x in db_batch):,.2f}")
    print(f"Diff: Rp {sum(x[2] for x in batch_excel) - sum(float(x['idr_amt']) for x in db_batch):,.2f}")

    matched_db = set()
    unmatched_ex = []
    for r, desc, amt in batch_excel:
        found = False
        for idx, d in enumerate(db_batch):
            if idx in matched_db:
                continue
            if abs(float(d['idr_amt']) - float(amt)) < 0.01:
                matched_db.add(idx)
                found = True
                break
        if not found:
            unmatched_ex.append((r, desc, amt))

    print(f"Unmatched in Excel Batch 2026: {len(unmatched_ex)}")
    for r, desc, amt in unmatched_ex:
        print(f"  Row {r}: {desc[:40]} | Rp {amt:,.2f}")

    unmatched_db = [db_batch[i] for i in range(len(db_batch)) if i not in matched_db]
    print(f"Unmatched in DB Batch 2026: {len(unmatched_db)}")
    for d in unmatched_db:
        print(f"  ID {d['id']}: {d['description'][:40]} | Rp {d['idr_amt']:,.2f}")

audit_year_2024()
audit_year_2025()
audit_year_2026()
conn.close()


