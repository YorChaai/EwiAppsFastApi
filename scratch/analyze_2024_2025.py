import psycopg2
from psycopg2.extras import DictCursor
import openpyxl
import datetime as dt

db_url = "postgresql://postgres:yorchai12@localhost:5432/miniproject_revisi_db"
conn = psycopg2.connect(db_url)
cur = conn.cursor(cursor_factory=DictCursor)

excel_path = r"D:\Aplikasi\EWIApps\data\excel\20250427_EWI Financial-Repport_2024-2026 (version 2).xlsx"
wb_val = openpyxl.load_workbook(excel_path, data_only=True)
wb_form = openpyxl.load_workbook(excel_path, data_only=False)

def analyze_year(year, sheet_name, r_single_start, r_single_end, r_batch_start, r_batch_end):
    print(f"\n=======================================================")
    print(f"             ANALISIS TAHUN {year}: {sheet_name}")
    print(f"=======================================================")
    ws_v = wb_val[sheet_name]
    ws_f = wb_form[sheet_name]

    # 1. Single analysis in Excel
    single_valid = []
    single_skipped = []
    for r in range(r_single_start, r_single_end + 1):
        val_b = ws_v.cell(r, 2).value
        val_d = ws_v.cell(r, 4).value
        amounts = []
        for c in range(5, 30):
            v = ws_v.cell(r, c).value
            if isinstance(v, (int, float)) and v != 0:
                amounts.append((c, v))
        desc_str = str(val_d or "").strip()
        all_empty = all(ws_v.cell(r, c).value is None for c in range(1, 30))

        if all_empty:
            single_skipped.append((r, "EMPTY ROW", ""))
        elif "TOTAL" in desc_str.upper() or "SUB TOTAL" in desc_str.upper():
            single_skipped.append((r, "SUBTOTAL", desc_str))
        elif not amounts and not isinstance(val_b, (dt.date, dt.datetime)):
            single_skipped.append((r, "HEADER / LABEL", desc_str or str(ws_v.cell(r, 2).value)))
        else:
            single_valid.append((r, val_b, desc_str, amounts))

    print(f"--- SINGLE MODE (Row {r_single_start} - {r_single_end}) ---")
    print(f"Total Baris Rentang: {r_single_end - r_single_start + 1}")
    print(f"Baris Valid Transaksi Single di Excel: {len(single_valid)}")
    print(f"Baris Dilewati (Header/Empty/Subtotal): {len(single_skipped)}")

    # 2. Batch analysis from Database
    cur.execute("""
        SELECT id, title, document_date, disbursed_amount, 
               (SELECT count(*) FROM expenses WHERE settlement_id = s.id) as item_count
        FROM settlements s
        WHERE report_year = %s AND settlement_type = 'batch'
        ORDER BY id ASC;
    """, (year,))
    batches = cur.fetchall()
    print(f"\n--- BATCH MODE DI DATABASE TAHUN {year} ({len(batches)} Batch) ---")
    tot_items = sum(b['item_count'] for b in batches)
    tot_amt = sum(b['disbursed_amount'] or 0 for b in batches)
    print(f"Total Batch Settlements: {len(batches)}")
    print(f"Total Item Pengeluaran Batch: {tot_items}")
    print(f"Total Nominal Batch: Rp {tot_amt:,.2f}")
    for idx, b in enumerate(batches, 1):
        print(f"  {idx:2d}. [ID {b['id']:4d}] Date: {b['document_date']} | Items: {b['item_count']:2d} | Amount: Rp {b['disbursed_amount']:>15,.2f} | Title: {b['title'][:55]}")

    # Single from DB
    cur.execute("SELECT count(*), sum(disbursed_amount) FROM settlements WHERE report_year = %s AND settlement_type = 'single';", (year,))
    s_cnt, s_amt = cur.fetchone()
    print(f"\n--- SINGLE MODE DI DATABASE TAHUN {year} ---")
    print(f"Total Single Settlements: {s_cnt}")
    print(f"Total Nominal Single: Rp {s_amt:,.2f}")

    print(f"\n--- TOTAL TAHUN {year} DI DATABASE ---")
    print(f"Total Settlement : {s_cnt + len(batches)} (Single: {s_cnt}, Batch: {len(batches)})")
    print(f"Total Nominal    : Rp {(s_amt or 0) + tot_amt:,.2f}")

# Run 2024
analyze_year(2024, "Rev VS Ops cost th 2024", 71, 263, 264, 864)

# Run 2025
analyze_year(2025, "Rev VS Ops cost th 2025_Final", 162, 516, 517, 1307)

conn.close()
