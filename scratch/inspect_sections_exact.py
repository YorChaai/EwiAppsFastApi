import openpyxl

wb_val = openpyxl.load_workbook('data/excel/20250427_EWI Financial-Repport_2024-2026 (version 2).xlsx', data_only=True)
wb_form = openpyxl.load_workbook('data/excel/20250427_EWI Financial-Repport_2024-2026 (version 2).xlsx', data_only=False)

def inspect_year(yr, sname, rev_range, tax_range, bank_range, single_range, batch_range):
    ws_v = wb_val[sname]
    ws_f = wb_form[sname]
    print(f"\n========================================================")
    print(f"   DETAILED INSPECTION: {yr} ({sname})")
    print(f"========================================================")
    
    # 1. REVENUE
    print(f"\n--- 1. REVENUE (Rows {rev_range[0]} - {rev_range[1]}) ---")
    rev_items = []
    rev_total_cell = None
    for r in range(rev_range[0], rev_range[1] + 1):
        vals = [(c, ws_v.cell(r, c).value) for c in range(1, 20) if ws_v.cell(r, c).value is not None]
        desc = str(ws_v.cell(r, 4).value or ws_v.cell(r, 3).value or '').strip()
        if 'TOTAL' in desc.upper() or 'TOTAL' in str(ws_v.cell(r, 2).value or '').upper():
            rev_total_cell = (r, vals)
        elif vals:
            rev_items.append((r, vals))
    print(f"  Total Data Rows: {len(rev_items)}")
    print(f"  Total Row: {rev_total_cell}")

    # 2. TAX
    print(f"\n--- 2. TAX (Rows {tax_range[0]} - {tax_range[1]}) ---")
    tax_items = []
    tax_total_cell = None
    for r in range(tax_range[0], tax_range[1] + 1):
        vals = [(c, ws_v.cell(r, c).value) for c in range(1, 20) if ws_v.cell(r, c).value is not None]
        desc = str(ws_v.cell(r, 4).value or ws_v.cell(r, 3).value or '').strip()
        if 'TOTAL' in desc.upper() or 'TOTAL' in str(ws_v.cell(r, 2).value or '').upper():
            tax_total_cell = (r, vals)
        elif vals:
            tax_items.append((r, vals))
    print(f"  Total Data Rows: {len(tax_items)}")
    print(f"  Total Row: {tax_total_cell}")

    # 3. BANK STATUS
    print(f"\n--- 3. BANK STATUS (Rows {bank_range[0]} - {bank_range[1]}) ---")
    bank_items = []
    bank_total_cell = None
    for r in range(bank_range[0], bank_range[1] + 1):
        vals = [(c, ws_v.cell(r, c).value) for c in range(1, 20) if ws_v.cell(r, c).value is not None]
        desc = str(ws_v.cell(r, 4).value or ws_v.cell(r, 3).value or '').strip()
        if 'TOTAL' in desc.upper() or 'TOTAL' in str(ws_v.cell(r, 2).value or '').upper():
            bank_total_cell = (r, vals)
        elif vals:
            bank_items.append((r, vals))
    print(f"  Total Data Rows: {len(bank_items)}")
    print(f"  Total Row: {bank_total_cell}")

inspect_year(2024, 'Rev VS Ops cost th 2024', (6, 22), (24, 37), (39, 68), (71, 263), (264, 861))
inspect_year(2025, 'Rev VS Ops cost th 2025_Final', (6, 46), (48, 82), (84, 159), (162, 516), (517, 1305))
inspect_year(2026, 'Rev VS Ops cost th 2026_', (6, 63), (65, 94), (96, 173), (176, 390), (391, 830))
