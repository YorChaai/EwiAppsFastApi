import openpyxl

excel_path = r"D:\Aplikasi\EWIApps\data\excel\20250427_EWI Financial-Repport_2024-2026 (version 2).xlsx"
sheet_name = "Rev VS Ops cost th 2026_"

wb_val = openpyxl.load_workbook(excel_path, data_only=True)
ws_v = wb_val[sheet_name]

# Check totals in Excel
for r in range(170, 835):
    for c in range(1, 10):
        v = str(ws_v.cell(r, c).value or "")
        if "TOTAL" in v.upper():
            vals = [(col, ws_v.cell(r, col).value) for col in range(5, 20) if ws_v.cell(r, col).value]
            print(f"Row {r:3d} Col {c}: {v[:40]:40s} | Values: {vals}")
