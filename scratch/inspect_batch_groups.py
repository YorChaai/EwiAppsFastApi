import openpyxl
import re
import datetime as dt

excel_path = r"D:\Aplikasi\EWIApps\data\excel\20250427_EWI Financial-Repport_2024-2026 (version 2).xlsx"
sheet_name = "Rev VS Ops cost th 2026_"

wb_val = openpyxl.load_workbook(excel_path, data_only=True)
wb_form = openpyxl.load_workbook(excel_path, data_only=False)

ws_v = wb_val[sheet_name]
ws_f = wb_form[sheet_name]

# Run the exact v14 detection logic for Batch groups on rows 391 - 830
from backend.scripts.generate_hybrid_migration_v14 import (
    is_divider_row, is_bank_header_row, _is_subcategory_cell, detect_cat, detect_subheader,
    to_num_cell, is_numeric_value, EXP_CAT_MAP
)

batch_groups = []
current_group = None
batch_items = []
unassigned_expenses = []

# Let's see how v14 processes row 391 to 830
is_in_batch_section = False
expecting_batch_title = False
col_curr = 7
col_rate = 8

for r in range(391, 831):
    val_a = ws_v.cell(r, 1).value
    val_b = ws_f.cell(r, 2).value
    val_c = ws_f.cell(r, 3).value
    val_d = ws_f.cell(r, 4).value

    cell_a = ws_f.cell(r, 1)
    cell_b = ws_f.cell(r, 2)
    cell_c = ws_f.cell(r, 3)
    cell_d = ws_f.cell(r, 4)

    val_a_str = str(val_a or "").upper().strip()
    val_b_str = str(val_b or "").upper().strip()
    val_c_str = str(val_c or "").upper().strip()
    val_d_str = str(val_d or "").upper().strip()

    # 0. Skip Noise Headers
    if any(h in val_d_str or h in val_c_str for h in ["DETAIL/DESCRIPTION", "ACTIVITY", "DESCRIPTION"]) or any(h in val_b_str or h in val_a_str for h in ["DATE", "#"]):
        continue

    # 1. Divider
    if is_divider_row(val_b_str, cell_b, val_d_str, cell_d) or (val_c_str in ["OPERATION COST AND OFFICE", "EXPENSES REPORT"]) or (val_d_str in ["OPERATION COST AND OFFICE", "EXPENSES REPORT"]):
        is_in_batch_section = True
        continue

    has_date_meta = isinstance(val_a, (dt.date, dt.datetime)) or isinstance(val_b, (dt.date, dt.datetime))
    is_manual_group = any(kw in val_a_str or kw in val_b_str or kw in val_d_str for kw in ["SUBJECT", "EXPENSE#"])
    
    is_subcat_for_check = (
        (val_d and _is_subcategory_cell(cell_d))
        or (val_b and _is_subcategory_cell(cell_b))
        or (val_a and _is_subcategory_cell(cell_a))
        or (val_c and _is_subcategory_cell(cell_c))
    )
    is_solid_row = (
        (cell_a.fill.patternType == "solid")
        or (cell_b.fill.patternType == "solid")
        or (cell_c.fill.patternType == "solid")
        or (cell_d.fill.patternType == "solid")
    ) and not is_subcat_for_check

    if not has_date_meta:
        matched_title = None
        if is_manual_group or expecting_batch_title:
            for cand in [val_c, val_d, val_b, val_a]:
                if cand:
                    t = str(cand).strip()
                    t_clean = re.sub(r"(?i)subject|expense#|:", "", t).strip()
                    if t_clean:
                        matched_title = t_clean
                        break
            if is_manual_group and not matched_title:
                expecting_batch_title = True
                continue
            expecting_batch_title = False
        elif is_solid_row and not is_subcat_for_check:
            for cand in [val_c, val_d, val_b, val_a]:
                if cand:
                    t = str(cand).strip()
                    t_clean = re.sub(r"(?i)subject|expense#|:", "", t).strip()
                    if t_clean:
                        matched_title = t_clean
                        break
        if matched_title:
            if current_group:
                batch_groups.append(current_group)
            current_group = {"row": r, "title": matched_title, "items": []}
            continue

    # 3. Subcategory
    is_subcat = (
        (val_d and _is_subcategory_cell(cell_d))
        or (val_b and _is_subcategory_cell(cell_b))
        or (val_a and _is_subcategory_cell(cell_a))
        or (val_c and _is_subcategory_cell(cell_c))
    )
    has_curr_meta = (ws_f.cell(r, col_curr).value is not None) if col_curr != -1 else False

    if is_subcat and not (has_date_meta and has_curr_meta):
        continue

    desc_for_check = str(val_d or val_c or "").strip().upper()
    if desc_for_check in ["TOTAL", "SUBTOTAL", "-", ""]:
        continue

    # Check item
    found_in_col = False
    col_amt = 0
    for col, c_id in EXP_CAT_MAP.items():
        v_tmp = to_num_cell(ws_f, ws_v, r, col)
        if v_tmp != 0:
            col_amt = v_tmp
            found_in_col = True
            break
    if not found_in_col:
        for fallback_col in [5, 6, 9, 4]:
            v_tmp = to_num_cell(ws_f, ws_v, r, fallback_col)
            if v_tmp != 0:
                col_amt = v_tmp
                found_in_col = True
                break

    if found_in_col:
        desc = str(val_c if is_numeric_value(val_d) else (val_d or "-")).strip()
        date_v = val_a if isinstance(val_a, (dt.date, dt.datetime)) else (val_b if isinstance(val_b, (dt.date, dt.datetime)) else None)
        item_info = {"row": r, "date": str(date_v), "desc": desc, "amount": col_amt}
        if current_group:
            current_group["items"].append(item_info)
        else:
            unassigned_expenses.append(item_info)
        batch_items.append(item_info)

if current_group:
    batch_groups.append(current_group)

print(f"Total Batch Groups Terdeteksi: {len(batch_groups)}")
print(f"Total Items dalam Batch: {len(batch_items)}")
print(f"Unassigned Items (di luar grup): {len(unassigned_expenses)}")

total_batch_amount = 0
for idx, g in enumerate(batch_groups, 1):
    g_total = sum(it["amount"] for it in g["items"])
    total_batch_amount += g_total
    print(f"{idx:2d}. [Row {g['row']:3d}] {g['title'][:55]:55s} -> {len(g['items']):2d} items | Total: Rp {g_total:,.0f}")

print(f"\nTotal Nominal Seluruh Batch: Rp {total_batch_amount:,.0f}")
